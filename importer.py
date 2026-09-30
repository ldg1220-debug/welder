"""기존 엑셀('1.Welder list', '2.자격 시험 기록서') -> SQLite 임포트.

사용: python importer.py <엑셀파일.xlsx> [--reset]
개인정보가 포함되므로 엑셀/DB 파일은 git에 올리지 않는다(.gitignore).
"""
from __future__ import annotations

import re
import sys
from datetime import date, datetime

import openpyxl

import database
import logic

SHEET_WELDERS = "1.Welder list"
SHEET_TESTS = "2.자격 시험 기록서"

PROCESS_MAP = {"GMAW": "135 MAG", "MAG": "135 MAG", "GTAW": "141 TIG", "TIG": "141 TIG", "LBW": "521 LBW"}


def norm_result(v) -> str | None:
    """● -> Pass, fail/FAIL/X -> Fail, 대기·의뢰·빈칸 등 -> None(미확정)."""
    if v is None:
        return None
    s = str(v).strip()
    if s == "●":
        return logic.PASS
    if s.lower() in ("fail", "x"):
        return logic.FAIL
    return None


def parse_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})\.?\s*", v)
        if m:
            try:
                return date(*map(int, m.groups()))
            except ValueError:
                return None
    return None


def parse_birth(v) -> str | None:
    """770715 / '80.01.01' -> ISO. 2자리 연도: 30 초과는 19xx, 이하는 20xx."""
    if v is None:
        return None
    digits = re.sub(r"\D", "", str(v))
    if len(digits) != 6:
        return None
    yy, mm, dd = int(digits[:2]), int(digits[2:4]), int(digits[4:])
    try:
        return date(1900 + yy if yy > 30 else 2000 + yy, mm, dd).isoformat()
    except ValueError:
        return None


def parse_thickness(v) -> float | None:
    m = re.search(r"\d+(\.\d+)?", str(v)) if v is not None else None
    return float(m.group()) if m else None


def norm_process(v) -> str | None:
    if v is None:
        return None
    key = re.split(r"[\s()]+", str(v).strip().upper())[0]
    return PROCESS_MAP.get(key, str(v).strip().replace("\n", " "))


def norm_backing(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip().upper()
    return {"NB": "nb", "MB": "mb", "CERAMIC": "Ceramic", "N/A": "N/A"}.get(s, str(v).strip())


def norm_filler(v) -> str | None:
    return re.sub(r"^FM\s*(\d)", r"FM \1", str(v).strip()) if v else None


def import_workbook(conn, path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    stats = {"welders": 0, "tests": 0, "skipped": [], "auto_welders": 0, "status_mismatch": 0}

    # 1) 용접사
    for r in wb[SHEET_WELDERS].iter_rows(min_row=5, values_only=True):
        if not isinstance(r[0], int) or not r[5]:
            continue
        status = r[9] if r[9] in ("재직", "퇴사") else "미확인"
        conn.execute(
            "INSERT OR IGNORE INTO welders (welder_no,name,birth_date,company,role,status) VALUES (?,?,?,?,?,?)",
            (str(r[5]).strip(), str(r[3]).strip(), parse_birth(r[4]), r[2], r[8], status),
        )
        stats["welders"] += 1

    ids = dict(conn.execute("SELECT welder_no, id FROM welders").fetchall())

    # 2) 시험 기록 (BW: VT/RT/Bend, FW: VT/Macro)
    for n, r in enumerate(wb[SHEET_TESTS].iter_rows(min_row=6, values_only=True), start=6):
        if not isinstance(r[0], int):
            continue
        no, wdate = (str(r[6]).strip() if r[6] else None), parse_date(r[1])
        joint = str(r[8]).strip().upper() if r[8] else None
        if not no or not wdate or joint not in ("BW", "FW"):
            stats["skipped"].append((n, "필수값(용접사No/시험일/시편형태) 누락"))
            continue
        if no not in ids:  # 명단에 없는 용접사는 시험기록 인적사항으로 최소 등록
            cur = conn.execute(
                "INSERT INTO welders (welder_no,name,birth_date,company,role,status) VALUES (?,?,?,?,?,?)",
                (no, str(r[4]).strip(), parse_birth(r[5]), r[3], None, "미확인"))
            ids[no] = cur.lastrowid
            stats["auto_welders"] += 1
        vt, rt = norm_result(r[14]), norm_result(r[15])
        mb = norm_result(r[16] if joint == "BW" else r[17])
        status = logic.judge(joint, vt, rt, mb)
        sheet_pass = "합격" if r[18] == "합격" else None
        if sheet_pass and status != "합격":
            stats["status_mismatch"] += 1
        cur = conn.execute(
            """INSERT INTO test_records (welder_id,project_name,welding_date,process,joint_type,thickness,
               position,filler_material,backing) VALUES (?,?,?,?,?,?,?,?,?)""",
            (ids[no], r[2], wdate.isoformat(), norm_process(r[7]), joint, parse_thickness(r[9]),
             str(r[10]).strip() if r[10] else None, norm_filler(r[11]), norm_backing(r[13])))
        conn.execute(
            """INSERT INTO test_results_and_certs (test_id,vt_result,rt_ut_result,macro_bend_result,
               final_status,issue_date,expire_date,renewal_count) VALUES (?,?,?,?,?,?,?,0)""",
            (cur.lastrowid, vt, rt, mb, status, None, logic.calc_expire_date(wdate).isoformat()))
        stats["tests"] += 1
    conn.commit()
    return stats


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    conn = database.get_conn()
    if "--reset" in sys.argv:
        conn.executescript("DROP TABLE IF EXISTS test_results_and_certs; DROP TABLE IF EXISTS test_records; DROP TABLE IF EXISTS welders;")
    database.init_db(conn)
    print(import_workbook(conn, sys.argv[1]))
